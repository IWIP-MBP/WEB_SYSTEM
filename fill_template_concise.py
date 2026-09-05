import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Set stdout encoding
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PIC_DIR = r"d:\WEB_SYSTEM\PIC"
TEMPLATE_PATH = os.path.join(PIC_DIR, "模板.xlsx")
OUTPUT_PATH = os.path.join(PIC_DIR, "后勤三部精益管理降本增效提案汇总表.xlsx")

def fill_excel_concise():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "精益管理降本增效提案表"

    headers = [
        "序号", 
        "车间", 
        "提案名称", 
        "提案分类", 
        "提案人/负责人", 
        "改善前现状及存在的问题", 
        "改善思路与方案", 
        "改善预期目标", 
        "改善前照片"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark blue
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border_thin

    # Simplified, non-technical, single-line proposals
    proposals = [
        {
            "seq": 1,
            "workshop": "后勤三部",
            "name": "电子花名册与极速检索",
            "category": "流程优化",
            "owner": "张金刚",
            "problem": "传统纸质与Excel台账分散混乱，查找员工信息耗时费力且易错。",
            "solution": "建立集中化电子花名册，实现中印双语姓名、车间与状态一键极速查询。",
            "target": "人事数据秒级检索，大幅提升查找效率，彻底消除录入与查询错漏。",
            "photo": ""
        },
        {
            "seq": 2,
            "workshop": "后勤三部",
            "name": "用工趋势与流动率分析",
            "category": "用工管理",
            "owner": "人力组",
            "problem": "缺乏用工流动预警，离职人员较多时无法及时补员，影响生产调度。",
            "solution": "建立入职与离职双趋势同屏对比看板，实时监控各车间人员流动与定编情况。",
            "target": "直观把控人员流动率，提前预判用工缺口，平抑用工波动并降低招聘成本。",
            "photo": ""
        },
        {
            "seq": 3,
            "workshop": "后勤三部",
            "name": "组织架构与本土化率监控",
            "category": "风险防控",
            "owner": "合规组",
            "problem": "海外印尼籍与中籍员工比例统计滞后，难以实时掌握合规用工情况。",
            "solution": "设立动态组织架构树状图，自动计算并实时展示印尼籍员工本土化率。",
            "target": "实时掌握跨国用工比例，确保完全符合当地法规要求，规避合规诉讼风险。",
            "photo": ""
        },
        {
            "seq": 4,
            "workshop": "后勤三部",
            "name": "劳保用品生命周期与防浪费",
            "category": "物资降本",
            "owner": "劳保组",
            "problem": "劳保领用缺乏上限控制与到期提醒，容易出现超领、重复领用与物资浪费。",
            "solution": "建立劳保发放与领用履历台账，设置换发周期并在到期前自动提醒。",
            "target": "杜绝无序超领与重复领用，预计降低劳保采购损耗20%，提升物资周转率。",
            "photo": ""
        },
        {
            "seq": 5,
            "workshop": "后勤三部",
            "name": "考勤日志极速自动对账",
            "category": "效率提升",
            "owner": "考勤组",
            "problem": "每月人工核对数千名员工考勤打卡与排休耗时数天，且易看错漏看引发争议。",
            "solution": "导入原始打卡记录与排休模板，自动进行考勤比对并自动标记异常情况。",
            "target": "考勤核算由3天缩短至15分钟，核算准确率达100%，消除薪酬争议。",
            "photo": ""
        },
        {
            "seq": 6,
            "workshop": "后勤三部",
            "name": "操作留痕与数据自动灾备",
            "category": "安全保障",
            "owner": "运维组",
            "problem": "缺少操作记录无法追溯责任，且无定时自动备份，存在数据损坏丢失风险。",
            "solution": "建立系统关键操作日志留痕机制，并设置每日自动备份与恢复功能。",
            "target": "关键操作全流程可追溯，确保数据安全不丢失，零额外运维资金投入。",
            "photo": ""
        }
    ]

    font_data = Font(name="Microsoft YaHei", size=10)
    font_bold = Font(name="Microsoft YaHei", size=10, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    row_start = 2
    for idx, p in enumerate(proposals):
        row_num = row_start + idx
        ws.row_dimensions[row_num].height = 24 # Standard clean row height for single line

        ws.cell(row=row_num, column=1, value=p["seq"]).alignment = align_center
        ws.cell(row=row_num, column=2, value=p["workshop"]).alignment = align_center
        ws.cell(row=row_num, column=3, value=p["name"]).alignment = align_left
        ws.cell(row=row_num, column=4, value=p["category"]).alignment = align_center
        ws.cell(row=row_num, column=5, value=p["owner"]).alignment = align_center
        ws.cell(row=row_num, column=6, value=p["problem"]).alignment = align_left
        ws.cell(row=row_num, column=7, value=p["solution"]).alignment = align_left
        ws.cell(row=row_num, column=8, value=p["target"]).alignment = align_left
        ws.cell(row=row_num, column=9, value=p["photo"]).alignment = align_center

        for col_idx in range(1, 10):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = font_data
            cell.border = border_thin
            if idx % 2 == 1:
                cell.fill = fill_zebra

        ws.cell(row=row_num, column=3).font = font_bold

    col_widths = {
        "A": 6,
        "B": 12,
        "C": 24,
        "D": 14,
        "E": 12,
        "F": 45,
        "G": 48,
        "H": 48,
        "I": 12
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(TEMPLATE_PATH)
    print(f"Updated template Excel at: {TEMPLATE_PATH}")
    wb.save(OUTPUT_PATH)
    print(f"Updated summary Excel at: {OUTPUT_PATH}")

if __name__ == "__main__":
    fill_excel_concise()
