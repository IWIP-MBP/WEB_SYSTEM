import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title="💡 操作提示", border_color="0F766E", bg_color="F0FDFA"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_t = p.add_run(f"{title}：")
    run_t.bold = True
    run_t.font.name = "微软雅黑"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = RGBColor(15, 118, 110)
    
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_image_with_caption(doc, img_path, caption_text):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(f"图：{caption_text}")
        run_cap.font.name = "微软雅黑"
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(100, 116, 139)

def build_docx_manual():
    doc = docx.Document()
    
    # Page setup - Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("后勤三部人事与劳保管理系统 v2.0 - 用户操作说明书")
        hrun.font.name = "微软雅黑"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("机密资料 · 仅供后勤三部内部使用")
        frun.font.name = "微软雅黑"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = '微软雅黑'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(51, 65, 85)

    # ------------------ 封面 ------------------
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(36)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("后勤三部人事与劳保管理系统")
    r_title.font.size = Pt(26)
    r_title.bold = True
    r_title.font.color.rgb = RGBColor(15, 118, 110)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("用户操作与系统使用说明书（v2.0 实际运行版）")
    r_sub.font.size = Pt(15)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    p_sub.paragraph_format.space_after = Pt(40)

    # Key metadata table on cover page
    table_meta = doc.add_table(rows=5, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("系统名称", "后勤三部人事与劳保管理系统 (HR & PPE System)"),
        ("适用部门", "印尼园区后勤三部（炼铁、烧结、动力等车间及班组）"),
        ("核心功能", "员工花名册管理、双语支持、劳保发放全追溯、考勤自动对账、数据容灾备份"),
        ("系统架构", "Streamlit 2.0 + FastAPI + PostgreSQL 15 (Docker 容器化部署)"),
        ("编制单位 / 作者", "后勤三部 张金刚 | 2026年8月最新版")
    ]
    for idx, (k, v) in enumerate(meta_data):
        cell_k = table_meta.cell(idx, 0)
        cell_v = table_meta.cell(idx, 1)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_margins(cell_k, top=100, bottom=100, left=150, right=150)
        set_cell_margins(cell_v, top=100, bottom=100, left=150, right=150)
        
        rk = cell_k.paragraphs[0].add_run(k)
        rk.bold = True
        rk.font.size = Pt(10)
        rk.font.color.rgb = RGBColor(30, 41, 59)
        
        rv = cell_v.paragraphs[0].add_run(v)
        rv.font.size = Pt(10)
        rv.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # Helper function for adding headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "微软雅黑"
        run.font.size = Pt(16)
        run.bold = True
        run.font.color.rgb = RGBColor(15, 118, 110)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "微软雅黑"
        run.font.size = Pt(12.5)
        run.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(text)
        run.font.name = "微软雅黑"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # ------------------ 正文内容 ------------------
    add_h1("一、 系统概述与总体架构")
    add_p("后勤三部人事与劳保管理系统（v2.0）是专门面向中印尼跨国工厂复杂作业环境打造的综合数字化管理平台。本系统旨在解决传统纸质档案查找困难、劳保用品领用无追溯、中印尼员工跨语言沟通障碍及考勤排休手动核算效率低下等核心痛点。")
    add_p("系统基于现代精益管理理念开发，具备以下核心亮点：")
    add_p("1. 中英/印尼双语实时切换：满足中方管理人员与印尼本地员工的无缝协作需求。\n"
          "2. 身份证件脱敏与安全控制：自动掩码脱敏敏感身份证号，防范隐私泄漏。\n"
          "3. 劳保全生命周期追溯与到期预警：精确跟踪手套、安全鞋、安全帽等物资的发放与使用周期，自动触发预警。\n"
          "4. 考勤排休智能转换：一键对账打卡日志与排休模板，处理跨零点打卡及复杂考勤。\n"
          "5. 自动化备份与灾备恢复：支持每日自动备份与一键秒级还原，保障核心数据绝对安全。")

    add_h1("二、 系统登录与安全认证")
    add_p("用户可通过局域网内部浏览器访问系统前端界面（默认端口 8501），系统提供严密的安全身份认证机制。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\01_系统登录界面.png", "系统登录与双语选择界面")
    add_h2("2.1 操作步骤")
    add_p("1. 打开 Edge / Chrome 浏览器，输入访问地址 `http://localhost:8501`（局域网用户请输入服务器 IP 地址）。\n"
          "2. 语言选择：点击界面上方【中文】或【Bahasa Indonesia】按钮，系统将即时切换全界面语言。\n"
          "3. 凭证录入：在用户名输入框填入账号（如 `admin`），在密码框填入对应密码，点击【登录】按钮。\n"
          "4. 快捷访问：登录页底栏集成了印尼语在线学习系统的跳转链接与版本说明。")
    add_callout(doc, "系统支持普通只读演示账号（如 viewer 角色）与超级管理员账号（admin 角色）的分级权限控制，演示账号可浏览所有模块但禁止修改与导出数据。", title="权限说明")

    add_h1("三、 数据看板与可视化决策")
    add_p("数据看板为管理者提供全盘人力资源与劳保状态的实时大盘视图。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\02_数据看板与双趋势图.png", "数据看板核心指标与双趋势分析图表")
    add_h2("3.1 核心指标与图表说明")
    add_p("• 人力核心指标卡：顶部实时统计在职员工总数、中方员工数、印尼本地员工数、车间班组覆盖数及本月流失率。\n"
          "• 本土化率分析：通过高对比度饼图直观展示印尼员工占比（如本土化率达 85.5%），助力工厂管理本土化推进。\n"
          "• 动态走势折线图：展示近 12 个月在职与离职人数变化曲线，帮助预判用工高峰与流失趋势。\n"
          "• 车间人员编制对比图：柱状图对比各车间实际到岗人数与编制目标。\n"
          "• 历史提醒挂件：侧边栏【📜 历史提醒】聚合展示未来 7 天内员工生日及劳保物品即将到期的人员名单。")

    add_h1("四、 组织架构图与本土化率")
    add_p("组织架构模块树状展现后勤三部下辖炼铁车间、烧结车间、动力车间等各个车间及班组的层级关系。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\04_组织架构与本土化率.png", "组织架构与班组本土化率统计视图")
    add_h2("4.1 功能特色")
    add_p("1. 层级结构展示：直观查看每个车间下属的各个生产班组与岗位配置。\n"
          "2. 班组本土化率核算：实时计算每个班组的中方与印尼籍员工比例。\n"
          "3. 架构数据导出：支持导出标准的 Excel 组织架构表，便于上报与汇总。")

    add_h1("五、 员工花名册全局管理")
    add_p("员工花名册是系统的核心基础数据库，管理全厂员工的全生命周期档案。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\03_员工花名册全局检索.png", "员工花名册全局检索与档案维护界面")
    add_h2("5.1 核心功能操作")
    add_p("• 即时模糊搜索：在顶部【实时搜索】框中输入工号（ID Nomor）或姓名（Nama），无需回车即可实时筛选表格。\n"
          "• 批量导入与模版：点击【批量导入员工】，可先下载标准 Excel 模版，填妥后上传一键导入大量人员数据。\n"
          "• 档案新增与修改：支持录入工号、姓名、国籍、车间、班组、岗位、入职日期、联系电话及身份证号。\n"
          "• 敏感数据脱敏：非高特权用户查看花名册时，身份证号自动显示为 `3201************12`，兼顾管理与隐私合规。")
    add_callout(doc, "修改员工工号时，系统会自动级联更新其历史异动记录、劳保领用日志及审计日志，确保数据完整性。", title="级联更新机制")

    add_h1("六、 离职名册与人员动态管理")
    add_p("离职名册负责妥善归档所有离职人员档案，并提供完善的档案复原机制。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\05_离职人员动态管理.png", "离职名册归档与一键复职操作界面")
    add_h2("6.1 离职与复职流程")
    add_p("1. 办理离职：在花名册中选中员工点击【办理离职】，填写离职日期与离职原因（如个人原因、合同到期），系统自动将其转入离职名册。\n"
          "2. 档案归档：离职名册完整保留员工离职前所在车间、工号、离职时间及经办人信息。\n"
          "3. 一键复职：若离职员工重新入职，管理员可在离职名册中点击【一键复职】，即可将其无缝还原回在职花名册并保留历史档案。")

    add_h1("七、 劳保用品全生命周期管理")
    add_p("劳保管理模块实现了后勤三部所有劳保用品（安全帽、防护服、安全鞋、手套、口罩等）的采购、库存、发放、领用及到期预警全流程数字化。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\06_劳保用品全生命周期管理.png", "劳保用品库存、领用发放与预警界面")
    add_h2("7.1 业务办理说明")
    add_p("• 劳保品类与库存维护：定义劳保物资规格、单价及标准使用周期（如安全鞋周期 180 天）。\n"
          "• 入库与出库登记：记录供应商入库及车间领用出库数量，实时更新库存余额。\n"
          "• 领用发放与对账：选择指定员工并发放劳保用品，系统自动根据上一次发放时间与标准周期推算下一次预计领用日期。\n"
          "• 到期预警与提醒：当员工劳保领用达到或超过周期时，主界面与任务栏弹出橙色预警提醒，支持批量导出领用签收单。")

    add_h1("八、 考勤排休自动化转换")
    add_p("针对工厂复杂的轮班与排休制度，系统提供了一键式打卡日志与排休模板智能对账转换功能。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\07_考勤自动对账与转换.png", "考勤原始打卡日志与排休模板智能转换界面")
    add_h2("8.1 转换操作步骤")
    add_p("1. 上传刷卡日志：在【上传考勤日志】区拖拽上传门禁机/打卡机导出的 Excel 明细表。\n"
          "2. 上传排休模板：在【上传排休模板】区上传标准排休安排 Excel 表。\n"
          "3. 一键转换：点击【🚀 开始考勤转换】，系统算法在后台自动完成零点跨天匹配、迟到早退判定及排休覆盖计算。\n"
          "4. 下载结果：转换完成后，直接点击【💾 下载考勤结果文件】获取标准对账分析表。")

    add_h1("九、 操作审计日志与安全管控")
    add_p("系统内置合规性审计日志模块，记录所有涉及数据写、改、删及配置变更的操作。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\08_系统审计日志与安全管控.png", "系统审计日志与实时操作跟踪界面")
    add_h2("9.1 审计追溯功能")
    add_p("• 详细记录内容：包括操作时间、操作员账号、客户端 IP 地址、操作类型（入职/修改/离职/异动）以及变更前与变更后的完整 JSON 差异。\n"
          "• 条件检索：支持按工号、姓名或时间范围快速检索特定操作记录。\n"
          "• 日志导出：支持将合规审计日志导出为标准 Excel 文件以备安全检查。")

    add_h1("十、 数据库自动化备份与容灾")
    add_p("为防止突发硬件故障或误操作造成数据丢失，系统配备了全自动数据库备份与秒级容灾还原机制。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\09_全自动备份与容灾管理.png", "自动化定时备份策略与快照还原界面")
    add_h2("10.1 备份与还原管理")
    add_p("1. 自动定时备份：管理员可设置每日定时备份时间（如凌晨 02:00）及快照保留天数（默认 7 天），过期备份自动清理。\n"
          "2. 立即备份：点击【立即备份】可实时创建全量 SQL 数据库快照。\n"
          "3. 一键秒级还原：选中历史备份快照点击【还原】，系统弹出二次确认窗口，确认后即可完成数据库全量还原。")

    add_h1("十一、 系统设置与参数维护")
    add_p("系统设置模块提供基础元数据、用户权限账号、企业 Logo 及登录页公告的全局配置。")
    add_image_with_caption(doc, r"d:\WEB_SYSTEM\PIC\10_系统设置与参数配置.png", "系统参数设置、用户管理与 Logo 配置界面")
    add_h2("11.1 配置项说明")
    add_p("• 元数据维护：动态维护车间列表、班组列表及员工国籍字典。\n"
          "• 用户与权限管理：支持管理员添加新用户账号，分配 `admin`（管理员）或 `viewer`（只读视角）角色，并支持重置密码。\n"
          "• 企业 Logo 自定义：上传企业专属 Logo 图片，系统顶部与任务栏图标自动同步替换。\n"
          "• 登录公告管理：可编辑登录页的中英/印尼双语欢迎词与通知公告。")

    # Save to docx file
    out_docx = r"d:\WEB_SYSTEM\PIC\后勤三部人事与劳保管理系统使用说明书.docx"
    doc.save(out_docx)
    print(f"Word manual created successfully at: {out_docx}")

def build_markdown_manual():
    md_content = """# 后勤三部人事与劳保管理系统 使用说明书 (v2.0 实际运行版)

**编制单位**：后勤三部 张金刚  
**适用对象**：后勤三部管理人员、HR 专员及车间主管  
**系统部署地址**：`http://localhost:8501`（局域网用户请访问服务器 IP）

---

## 一、 系统概述与总体架构

**后勤三部人事与劳保管理系统（v2.0）**是专门面向中印尼跨国工厂复杂作业环境打造的综合数字化管理平台。本系统旨在解决传统纸质档案查找困难、劳保用品领用无追溯、中印尼员工跨语言沟通障碍及考勤排休手动核算效率低下等核心痛点。

### 核心亮点
1. **中英/印尼双语实时切换**：满足中方管理人员与印尼本地员工的无缝协作需求。
2. **身份证件脱敏与安全控制**：自动掩码脱敏敏感身份证号，防范隐私泄漏。
3. **劳保全生命周期追溯与到期预警**：精确跟踪手套、安全鞋、安全帽等物资的发放与使用周期，自动触发预警。
4. **考勤排休智能转换**：一键对账打卡日志与排休模板，处理跨零点打卡及复杂考勤。
5. **自动化备份与灾备恢复**：支持每日自动备份与一键秒级还原，保障核心数据绝对安全。

---

## 二、 系统登录与安全认证

用户可通过局域网内部浏览器访问系统前端界面（默认端口 8501），系统提供严密的安全身份认证机制。

![系统登录与双语选择界面](01_系统登录界面.png)

### 操作步骤
1. **访问系统**：打开 Edge / Chrome 浏览器，输入访问地址 `http://localhost:8501`。
2. **语言切换**：点击界面上方【中文】或【Bahasa Indonesia】按钮，系统将即时切换全界面语言。
3. **凭证录入**：在用户名输入框填入账号（如 `admin`），在密码框填入对应密码，点击【登录】按钮。
4. **快捷访问**：登录页底栏集成了印尼语在线学习系统的跳转链接与版本说明。

> 💡 **权限说明**：系统支持普通只读演示账号（如 viewer 角色）与超级管理员账号（admin 角色）的分级权限控制，演示账号可浏览所有模块但禁止修改与导出数据。

---

## 三、 数据看板与可视化决策

数据看板为管理者提供全盘人力资源与劳保状态的实时大盘视图。

![数据看板核心指标与双趋势分析图表](02_数据看板与双趋势图.png)

### 核心指标与图表说明
- **人力核心指标卡**：顶部实时统计在职员工总数、中方员工数、印尼本地员工数、车间班组覆盖数及本月流失率。
- **本土化率分析**：通过高对比度饼图直观展示印尼员工占比（如本土化率达 85.5%），助力工厂管理本土化推进。
- **动态走势折线图**：展示近 12 个月在职与离职人数变化曲线，帮助预判用工高峰与流失趋势。
- **车间人员编制对比图**：柱状图对比各车间实际到岗人数与编制目标。
- **历史提醒挂件**：侧边栏【📜 历史提醒】聚合展示未来 7 天内员工生日及劳保物品即将到期的人员名单。

---

## 四、 组织架构图与本土化率

组织架构模块树状展现后勤三部下辖炼铁车间、烧结车间、动力车间等各个车间及班组的层级关系。

![组织架构与班组本土化率统计视图](04_组织架构与本土化率.png)

### 功能特色
1. **层级结构展示**：直观查看每个车间下属的各个生产班组与岗位配置。
2. **班组本土化率核算**：实时计算每个班组的中方与印尼籍员工比例。
3. **架构数据导出**：支持导出标准的 Excel 组织架构表，便于上报与汇总。

---

## 五、 员工花名册全局管理

员工花名册是系统的核心基础数据库，管理全厂员工的全生命周期档案。

![员工花名册全局检索与档案维护界面](03_员工花名册全局检索.png)

### 核心功能操作
- **即时模糊搜索**：在顶部【实时搜索】框中输入工号（ID Nomor）或姓名（Nama），无需回车即可实时筛选表格。
- **批量导入与模版**：点击【批量导入员工】，可先下载标准 Excel 模版，填妥后上传一键导入大量人员数据。
- **档案新增与修改**：支持录入工号、姓名、国籍、车间、班组、岗位、入职日期、联系电话及身份证号。
- **敏感数据脱敏**：非高特权用户查看花名册时，身份证号自动显示为 `3201************12`，兼顾管理与隐私合规。

> 💡 **级联更新机制**：修改员工工号时，系统会自动级联更新其历史异动记录、劳保领用日志及审计日志，确保数据完整性。

---

## 六、 离职名册与人员动态管理

离职名册负责妥善归档所有离职人员档案，并提供完善的档案复原机制。

![离职名册归档与一键复职操作界面](05_离职人员动态管理.png)

### 离职与复职流程
1. **办理离职**：在花名册中选中员工点击【办理离职】，填写离职日期与离职原因（如个人原因、合同到期），系统自动将其转入离职名册。
2. **档案归档**：离职名册完整保留员工离职前所在车间、工号、离职时间及经办人信息。
3. **一键复职**：若离职员工重新入职，管理员可在离职名册中点击【一键复职】，即可将其无缝还原回在职花名册并保留历史档案。

---

## 七、 劳保用品全生命周期管理

劳保管理模块实现了后勤三部所有劳保用品（安全帽、防护服、安全鞋、手套、口罩等）的采购、库存、发放、领用及到期预警全流程数字化。

![劳保用品库存、领用发放与预警界面](06_劳保用品全生命周期管理.png)

### 业务办理说明
- **劳保品类与库存维护**：定义劳保物资规格、单价及标准使用周期（如安全鞋周期 180 天）。
- **入库与出库登记**：记录供应商入库及车间领用出库数量，实时更新库存余额。
- **领用发放与对账**：选择指定员工并发放劳保用品，系统自动根据上一次发放时间与标准周期推算下一次预计领用日期。
- **到期预警与提醒**：当员工劳保领用达到或超过周期时，主界面与任务栏弹出橙色预警提醒，支持批量导出领用签收单。

---

## 八、 考勤排休自动化转换

针对工厂复杂的轮班与排休制度，系统提供了一键式打卡日志与排休模板智能对账转换功能。

![考勤原始打卡日志与排休模板智能转换界面](07_考勤自动对账与转换.png)

### 转换操作步骤
1. **上传刷卡日志**：在【上传考勤日志】区拖拽上传门禁机/打卡机导出的 Excel 明细表。
2. **上传排休模板**：在【上传排休模板】区上传标准排休安排 Excel 表。
3. **一键转换**：点击【🚀 开始考勤转换】，系统算法在后台自动完成零点跨天匹配、迟到早退判定及排休覆盖计算。
4. **下载结果**：转换完成后，直接点击【💾 下载考勤结果文件】获取标准对账分析表。

---

## 九、 操作审计日志与安全管控

系统内置合规性审计日志模块，记录所有涉及数据写、改、删及配置变更的操作。

![系统审计日志与实时操作跟踪界面](08_系统审计日志与安全管控.png)

### 审计追溯功能
- **详细记录内容**：包括操作时间、操作员账号、客户端 IP 地址、操作类型（入职/修改/离职/异动）以及变更前与变更后的完整 JSON 差异。
- **条件检索**：支持按工号、姓名或时间范围快速检索特定操作记录。
- **日志导出**：支持将合规审计日志导出为标准 Excel 文件以备安全检查。

---

## 十、 数据库自动化备份与容灾

为防止突发硬件故障或误操作造成数据丢失，系统配备了全自动数据库备份与秒级容灾还原机制。

![自动化定时备份策略与快照还原界面](09_全自动备份与容灾管理.png)

### 备份与还原管理
1. **自动定时备份**：管理员可设置每日定时备份时间（如凌晨 02:00）及快照保留天数（默认 7 天），过期备份自动清理。
2. **立即备份**：点击【立即备份】可实时创建全量 SQL 数据库快照。
3. **一键秒级还原**：选中历史备份快照点击【还原】，系统弹出二次确认窗口，确认后即可完成数据库全量还原。

---

## 十一、 系统设置与参数维护

系统设置模块提供基础元数据、用户权限账号、企业 Logo 及登录页公告的全局配置。

![系统参数设置、用户管理与 Logo 配置界面](10_系统设置与参数配置.png)

### 配置项说明
- **元数据维护**：动态维护车间列表、班组列表及员工国籍字典。
- **用户与权限管理**：支持管理员添加新用户账号，分配 `admin`（管理员）或 `viewer`（只读视角）角色，并支持重置密码。
- **企业 Logo 自定义**：上传企业专属 Logo 图片，系统顶部与任务栏图标自动同步替换。
- **登录公告管理**：可编辑登录页的中英/印尼双语欢迎词与通知公告。
"""
    out_md = r"d:\WEB_SYSTEM\PIC\后勤三部人事与劳保管理系统使用说明书.md"
    with open(out_md, "w", encoding="utf-8") as f:
        f.write(md_content)
    print(f"Markdown manual created successfully at: {out_md}")

if __name__ == "__main__":
    build_docx_manual()
    build_markdown_manual()
