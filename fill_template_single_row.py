import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

# Set stdout encoding
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

PIC_DIR = r"d:\WEB_SYSTEM\PIC"
TEMPLATE_PATH = os.path.join(PIC_DIR, "模板.xlsx")
OUTPUT_PATH = os.path.join(PIC_DIR, "后勤三部精益管理降本增效提案汇总表.xlsx")

def fill_excel_single_row():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "精益管理降本增效提案表"

    headers = [
        "序号", 
        "车间", 
        "提案名称", 
        "提案分类", 
        "提案人/负责人", 
        "改善前现状及存在的问题", 
        "改善思路与方案", 
        "改善预期目标", 
        "改善前照片"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark blue
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border_thin

    # Single comprehensive proposal row
    single_proposal = {
        "seq": 1,
        "workshop": "后勤三部",
        "name": "后勤数字化管理系统建设与降本增效",
        "category": "综合改善",
        "owner": "张金刚",
        "problem": "传统纸质与Excel台账分散混乱，查找慢、对账繁重易错，缺乏实时监控、到期预警与自动备份机制。",
        "solution": "建立集中化数字化平台，整合花名册检索、用工趋势预警、架构渲染、劳保到期提醒、考勤自动比对与每日定时灾备。",
        "target": "业务处理效率提升90%以上，数据100%精准无错漏，用工与本土化率实时掌控，审计留痕与自动灾备保障数据安全不丢失。",
        "photo": ""
    }

    font_data = Font(name="Microsoft YaHei", size=10)
    font_bold = Font(name="Microsoft YaHei", size=10, bold=True)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)

    row_num = 2
    ws.row_dimensions[row_num].height = 28

    ws.cell(row=row_num, column=1, value=single_proposal["seq"]).alignment = align_center
    ws.cell(row=row_num, column=2, value=single_proposal["workshop"]).alignment = align_center
    ws.cell(row=row_num, column=3, value=single_proposal["name"]).alignment = align_left
    ws.cell(row=row_num, column=4, value=single_proposal["category"]).alignment = align_center
    ws.cell(row=row_num, column=5, value=single_proposal["owner"]).alignment = align_center
    ws.cell(row=row_num, column=6, value=single_proposal["problem"]).alignment = align_left
    ws.cell(row=row_num, column=7, value=single_proposal["solution"]).alignment = align_left
    ws.cell(row=row_num, column=8, value=single_proposal["target"]).alignment = align_left
    ws.cell(row=row_num, column=9, value=single_proposal["photo"]).alignment = align_center

    for col_idx in range(1, 10):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.font = font_data
        cell.border = border_thin

    ws.cell(row=row_num, column=3).font = font_bold

    col_widths = {
        "A": 6,
        "B": 12,
        "C": 30,
        "D": 12,
        "E": 12,
        "F": 55,
        "G": 65,
        "H": 68,
        "I": 12
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    wb.save(TEMPLATE_PATH)
    print(f"Updated single-row template Excel at: {TEMPLATE_PATH}")
    wb.save(OUTPUT_PATH)
    print(f"Updated single-row summary Excel at: {OUTPUT_PATH}")

if __name__ == "__main__":
    fill_excel_single_row()
